#!/usr/bin/env python3
"""Normalize generated manuscript DOCX styles and accessibility metadata."""

from __future__ import annotations

import argparse
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

BLACK = RGBColor(0, 0, 0)
HEADER_FILL = "D9E2F3"
FIGURE_METADATA = [
    (
        "MetaboTyping Agentic Workbench evidence flow",
        "Four-stage diagram showing research inputs, explicit evidence models, deterministic decision "
        "modules, and reviewable outputs. Accepted, human-review, and non-combinable states remain "
        "separate, and a note states that the evaluated offline workflow does not call a language model.",
    ),
    (
        "Synthetic MoTrPAC-style volcano plot",
        "Volcano plot of six synthetic metabolites for endurance exercise 10 minutes after versus "
        "before exercise. Colors and symbols distinguish effect direction, FDR selection, and the "
        "effect threshold.",
    ),
    (
        "Synthetic reviewed metabolite effect heatmap",
        "Heatmap of synthetic log2 fold changes across two exercise contrasts. Symbols identify an "
        "unreviewed harmonization, an unreported value, and a missingness or limit-of-detection concern.",
    ),
]


def set_font_name(font, name: str) -> None:
    font.name = name
    rpr = font._element.get_or_add_rPr()
    rfonts = rpr.get_or_add_rFonts()
    for attribute in ["ascii", "hAnsi", "eastAsia", "cs"]:
        rfonts.set(qn(f"w:{attribute}"), name)


def set_run_font(run, name: str = "Arial") -> None:
    run.font.name = name
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ["ascii", "hAnsi", "eastAsia", "cs"]:
        rfonts.set(qn(f"w:{attribute}"), name)
    run.font.color.rgb = BLACK


def add_page_number(paragraph) -> None:
    paragraph.clear()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    for element in [begin, instruction, separate, value, end]:
        run._r.append(element)
    set_run_font(run)
    run.font.size = Pt(9)


def mark_header_row(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:tblHeader")) is None:
        row_properties.append(OxmlElement("w:tblHeader"))
    for cell in row.cells:
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell_properties = cell._tc.get_or_add_tcPr()
        shading = cell_properties.find(qn("w:shd"))
        if shading is None:
            shading = OxmlElement("w:shd")
            cell_properties.append(shading)
        shading.set(qn("w:fill"), HEADER_FILL)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True


def prevent_row_split(row) -> None:
    row_properties = row._tr.get_or_add_trPr()
    if row_properties.find(qn("w:cantSplit")) is None:
        row_properties.append(OxmlElement("w:cantSplit"))


def set_update_fields(doc) -> None:
    settings = doc.settings.element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    update_fields.set(qn("w:val"), "true")


def normalize(path: Path) -> None:
    doc = Document(path)

    for section in doc.sections:
        section.top_margin = Inches(0.85)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)
        section.footer_distance = Inches(0.35)
        add_page_number(section.footer.paragraphs[0])

    style_sizes = {
        "Normal": 10.5,
        "Title": 17,
        "Subtitle": 10.5,
        "Heading 1": 14,
        "Heading 2": 12,
        "Heading 3": 11,
        "Caption": 9,
    }
    for style in doc.styles:
        if style.type != WD_STYLE_TYPE.PARAGRAPH:
            continue
        set_font_name(style.font, "Arial")
        style.font.color.rgb = BLACK
        if style.name in style_sizes:
            style.font.size = Pt(style_sizes[style.name])
        if style.name.startswith("Heading"):
            style.font.bold = True

    normal = doc.styles["Normal"]
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08

    in_front_matter = True
    for paragraph in doc.paragraphs:
        style_name = paragraph.style.name
        if style_name.startswith("Heading"):
            in_front_matter = False
        for run in paragraph.runs:
            set_run_font(run)
        if style_name == "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(10)
        elif style_name == "Subtitle":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(3)
        elif style_name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True
            paragraph.paragraph_format.space_before = Pt(10)
            paragraph.paragraph_format.space_after = Pt(4)
        elif style_name == "Caption":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.keep_together = True
            paragraph.paragraph_format.keep_with_next = False
            paragraph.paragraph_format.space_after = Pt(6)

        if in_front_matter and style_name != "Title":
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(2.5)

        if paragraph._p.xpath(".//w:drawing"):
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.keep_with_next = True

    for index, shape in enumerate(doc.inline_shapes):
        title, description = FIGURE_METADATA[min(index, len(FIGURE_METADATA) - 1)]
        shape._inline.docPr.set("descr", description)
        shape._inline.docPr.set("title", title)

    for table in doc.tables:
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = True
        if table.rows:
            mark_header_row(table.rows[0])
        for row in table.rows:
            prevent_row_split(row)
            for cell in row.cells:
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_after = Pt(1.5)
                    for run in paragraph.runs:
                        set_run_font(run)
                        run.font.size = Pt(8.5)

    set_update_fields(doc)
    doc.core_properties.title = "MetaboTyping Agentic Workbench"
    doc.core_properties.subject = "JORS Software Metapaper draft"
    doc.core_properties.author = "Samuel Montalvo; Eric Leslie; Laurens van de Wiel; Matthew T. Wheeler"
    doc.core_properties.keywords = "metabolomics; harmonization; exercise; research software; provenance"
    doc.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("docx", type=Path)
    args = parser.parse_args()
    normalize(args.docx)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
